"""Create an evidence-based plan snapshot at the 24/09/2026 cutoff."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Ke_hoach_du_an_AI_dich_thuat_25-08_den_30-10-2026.xlsx"
OUTPUT = ROOT / "docs" / "Ke_hoach_du_an_AI_dich_thuat_hoan_thien_ky_thuat_25-09-2026.xlsx"


# percentage, evidence, PM assessment
PROGRESS = {
    "1.1": (1.00, "PROJECT_ARTIFACTS.md §1", "Đủ charter, RACI, DoD; cần chữ ký nếu dùng làm biên bản chính thức."),
    "1.2": (1.00, "inventory_baseline_v4/report.json", "Snapshot 227 source files, 45 route và báo cáo hiện trạng."),
    "1.3": (1.00, "PROJECT_ARTIFACTS.md §3–4", "Có ma trận frontend/admin và nợ kỹ thuật."),
    "1.4": (1.00, "PROJECT_ARTIFACTS.md §3; inventory report", "Đã ghi môi trường, Docker/env, GPU và phụ thuộc."),
    "1.5": (1.00, "PROJECT_ARTIFACTS.md §4", "Backlog/RTM ở mức quản lý; tiếp tục cập nhật sau cutoff."),
    "2.1": (1.00, "PROJECT_ARTIFACTS.md §1–2", "MVP, ngoài phạm vi và DoD đã được văn bản hóa."),
    "2.2": (1.00, "PROJECT_ARTIFACTS.md §5", "Luồng Admin/Reviewer/Superadmin được mô tả."),
    "2.3": (1.00, "PROJECT_ARTIFACTS.md §5; admin browser tests", "Wireframe văn bản + UI thực thi; responsive test đạt."),
    "2.4": (1.00, "benchmark_v2_500.csv; benchmark_v2_metadata.json; audio_benchmark_v2/", "Đủ 500 câu có version/hash và 100 audio sạch/nhiễu; corpus tổng hợp, không thay thế dữ liệu production."),
    "2.5": (1.00, "translation_benchmark_v2_summary.json; speech_benchmark_v2_summary.json", "Đã đo BLEU/chrF/token-F1/accuracy thuật ngữ/WER; kết quả dưới KPI là đầu vào tối ưu tiếp theo."),
    "2.6": (1.00, "full_pipeline_benchmark_summary.json; full_pipeline_profile.txt", "Đã đo warm local STT–NLLB–TTS p50/p95/p99, throughput và GPU; network/queue được công bố bằng 0 vì chạy local."),
    "2.7": (1.00, "PROJECT_ARTIFACTS.md §6; inventory route list", "Đã mô tả kiến trúc, REST/WS và mã lỗi."),
    "3.1": (1.00, "admin_build.log; admin_browser_tests.log", "Build đạt; responsive browser test đạt."),
    "3.2": (1.00, "backend_full_tests.log", "Auth/RBAC/API key/isolation test đạt."),
    "3.3": (1.00, "App.tsx; admin browser tests", "Login và protected routes hoạt động theo test."),
    "3.4": (1.00, "admin_api_large_dataset_benchmark.json; admin_api_profile.txt", "Đã benchmark 20.000 log và 5.001 user; có p50/p95/p99 và cProfile."),
    "3.5": (1.00, "Dashboard.tsx; Analytics.tsx; admin build/browser tests", "UI build và test đạt."),
    "3.6": (1.00, "admin_routes.py; admin extension tests", "CRUD user, role/password, reset session và create/revoke/rotate key có audit đạt kiểm thử."),
    "3.7": (1.00, "Users.tsx; ApiKeys.tsx; browser tests", "UI list/search/page/create/sửa role-mật khẩu/disable/reset/rotate/revoke đạt browser test."),
    "3.8": (1.00, "admin_routes.py; admin extension tests", "CRUD/import/export Dictionary, History, structured QA và audit đạt kiểm thử."),
    "3.9": (1.00, "Dictionary.tsx; History.tsx; QA.tsx; browser tests", "Luồng Dictionary/History/QA gồm scoring và export đã hoàn thiện kỹ thuật."),
    "3.10": (1.00, "ModelDeployment; /models/*; /audit; admin extension tests", "Settings whitelist, active/canary/promote/rollback có version và audit đạt kiểm thử."),
    "3.11": (1.00, "Settings.tsx; admin browser tests", "UI settings/model/audit có validate, canary, promote, rollback và bộ lọc audit."),
    "4.1": (1.00, "regression_results_v3.json; backend tests", "Các probe normalization/punctuation/hallucination đạt."),
    "4.2": (1.00, "translation_memory.py; translation quality tests", "Exact/fuzzy direction-aware, isolation, version/invalidation và ngưỡng chống match mơ hồ đạt test; KPI thuật ngữ vẫn là quality gap."),
    "4.3": (1.00, "translation_benchmark_v2_cases.csv; translation_benchmark_v2_summary.json", "Benchmark v2 đã chạy đủ 500/500 case trên NLLB local GPU."),
    "4.4": (1.00, "QA.tsx; QualityReview; admin extension tests", "Có sửa, adequacy/fluency, error taxonomy, critical flag, export và reviewer metadata."),
    "4.5": (0.50, "human_evaluation_v2_two_reviewers.csv", "Đã lấy mẫu 100 case × 2 reviewer và tạo phiếu; bắt buộc Hà Văn Đô và Nguyễn Quang Thọ chấm/ký, AI không được giả mạo."),
    "5.1": (1.00, "full_pipeline_profile.prof/.txt; admin_api_profile.prof/.txt", "Đã profile API dữ liệu lớn và pipeline model thật; có critical-path artifact tái lập."),
    "5.2": (1.00, "inference_ab_soak_summary.json", "A/B beam 1–3 và soak 200 lượt: beam 1 được chọn, không OOM, memory growth 0 MB."),
}


def main() -> None:
    wb = load_workbook(SOURCE)
    ws = wb["Ke_hoach_chi_tiet"]
    ws.cell(4, 17, "Bằng chứng kiểm tra ngày 25/09")
    ws.cell(4, 18, "Trạng thái sau bổ sung / Phần còn thiếu")
    fill = PatternFill("solid", fgColor="17365D")
    border = Border(*(Side(style="thin", color="B7C9D6") for _ in range(4)))
    for col in (17, 18):
        cell = ws.cell(4, col)
        cell.fill = fill
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in range(5, 5 + 61):
        wbs = str(ws.cell(row, 2).value)
        if wbs not in PROGRESS:
            continue
        percentage, evidence, assessment = PROGRESS[wbs]
        ws.cell(row, 10, "Hoàn thành" if percentage == 1 else "Chờ đánh giá độc lập")
        ws.cell(row, 16, percentage)
        ws.cell(row, 17, evidence)
        ws.cell(row, 18, assessment)
        for col in (17, 18):
            ws.cell(row, col).font = Font(name="Aptos", size=9)
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row, col).border = border
    # Các việc sau ngày chốt không phải là việc trễ; hiển thị đúng nghĩa lịch kế hoạch.
    for row in range(5, 5 + 61):
        wbs = str(ws.cell(row, 2).value)
        if wbs not in PROGRESS and ws.cell(row, 10).value == "Chưa bắt đầu":
            ws.cell(row, 10, "Chưa đến hạn")
    ws.column_dimensions[get_column_letter(17)].width = 38
    ws.column_dimensions[get_column_letter(18)].width = 55
    ws.auto_filter.ref = f"A4:R{4 + 61}"
    ws.print_area = f"A1:R{4 + 61}"

    milestones = wb["Moc_ban_giao"]
    milestones["F5"] = "Đạt"
    milestones["F6"] = "Đạt có điều kiện"

    kpi = wb["KPI_nghiem_thu"]
    actual = {
        "K01": "Benchmark v2 BLEU(smoothed) = 56,925; chrF = 74,867",
        "K03": "Benchmark v2: 650/875 = 74,29% — chưa đạt 95%",
        "K04": "Audio tổng hợp: WER sạch 39,03%; nhiễu 15 dB 40,33% — chưa đạt",
        "K06": "Warm local STT–NLLB–TTS p95 = 1.175,470 ms; không gồm network/queue",
        "K11": "Admin build đạt; browser test được chạy lại trong verification_summary.json",
        "K12": "Backend 121/121; admin browser 21/21; Flutter analyze đạt, không phát hiện lỗi",
    }
    for row in range(5, kpi.max_row + 1):
        code = kpi.cell(row, 1).value
        if code in actual:
            kpi.cell(row, 8, actual[code])

    overview = wb["Tong_quan"]
    overview["B3"] = "Hoàn thiện kỹ thuật sau đối chiếu mốc 24/09, cập nhật ngày 25/09/2026 — 1 hạng mục đánh giá người thật chờ hai thành viên ký."
    overview.merge_cells("B3:H3")
    overview["B3"].font = Font(name="Aptos", bold=True, color="C00000")
    overview["B3"].alignment = Alignment(wrap_text=True)
    wb.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Updated tasks: {len(PROGRESS)}; complete: {sum(value[0] == 1 for value in PROGRESS.values())}; partial: {sum(value[0] < 1 for value in PROGRESS.values())}")


if __name__ == "__main__":
    main()
